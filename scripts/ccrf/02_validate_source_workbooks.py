#!/usr/bin/env python3
from __future__ import annotations
import argparse
import json
from pathlib import Path
from openpyxl import load_workbook
from common import BUILD, VALIDATION, canonical_header_map, load_sources, sha256_file


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument('--input-dir', default=str(BUILD / 'ccrf_sources'))
    args = parser.parse_args()
    cfg = load_sources()
    src = Path(args.input_dir)
    report = {'ccrf_version': cfg['ccrf_version'], 'files': {}, 'passed': True}

    for key, source in cfg['sources'].items():
        path = src / source['filename']
        item = {'path': str(path), 'exists': path.exists()}

        if not path.exists():
            report['passed'] = False
            report['files'][key] = item
            continue

        item['sha256'] = sha256_file(path)
        item['hash_matches'] = item['sha256'] == source['sha256']

        workbook = load_workbook(path, read_only=True, data_only=False)
        sheet = workbook[workbook.sheetnames[0]]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        item['sheet'] = sheet.title
        item['header_map'] = canonical_header_map(headers)

        variable_index = item['header_map']['variable']
        data_rows = 0
        for row in sheet.iter_rows(min_row=2, values_only=True):
            value = row[variable_index] if variable_index < len(row) else None
            if value is not None and str(value).strip():
                data_rows += 1

        item['data_rows'] = data_rows
        item['expected_rows'] = source['expected_rows']
        item['row_count_matches'] = item['data_rows'] == source['expected_rows']
        item['passed'] = item['hash_matches'] and item['row_count_matches']

        report['passed'] = report['passed'] and item['passed']
        report['files'][key] = item

    VALIDATION.mkdir(exist_ok=True)
    (VALIDATION / 'ccrf_source_validation.json').write_text(
        json.dumps(report, indent=2) + '\n',
        encoding='utf-8',
    )
    print(json.dumps(report, indent=2))
    return 0 if report['passed'] else 1


if __name__ == '__main__':
    raise SystemExit(main())
