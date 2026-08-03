#!/usr/bin/env python3
from __future__ import annotations
import argparse
import json
from pathlib import Path
from openpyxl import load_workbook
from common import BUILD, DATA, canonical_header_map, clean, derive_section, bool_yn, load_sources, load_official_counts, parse_bound, parse_snapshot

INPUT_TYPES = {'text','notes','radio','dropdown','checkbox','slider','file'}


def parse_xlsx(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    headers = list(next(rows))
    header_map = canonical_header_map(headers)
    fields = []
    for idx, row in enumerate(rows, start=2):
        def get(key):
            return clean(row[header_map[key]]) if key in header_map and header_map[key] < len(row) else None
        variable = get('variable')
        if not variable:
            continue
        fields.append({
            'variable': variable,
            'form': get('form') or 'Unassigned',
            'sourceSectionHeader': get('sourceSectionHeader'),
            'type': get('type') or 'text',
            'label': get('label'),
            'choices': get('choices'),
            'note': get('note'),
            'validation': get('validation'),
            'min': parse_bound(row[header_map['min']]) if 'min' in header_map and header_map['min'] < len(row) else None,
            'max': parse_bound(row[header_map['max']]) if 'max' in header_map and header_map['max'] < len(row) else None,
            'identifier': bool_yn(get('identifier')),
            'branching': get('branching'),
            'required': bool_yn(get('required')),
            'alignment': get('alignment'),
            'questionNumber': get('questionNumber'),
            'matrixGroup': get('matrixGroup'),
            'matrixRanking': get('matrixRanking'),
            'annotation': get('annotation'),
            'sourceRow': idx,
        })
    return derive_section(fields)


def build_schema(groups, cfg, official):
    classes = {}
    for key, fields in groups.items():
        source = cfg['sources'][key]
        count_meta = official['classifications'][key]
        if count_meta['clinical_cdes'] + count_meta['participant_cdes'] != count_meta['official_cdes']:
            raise ValueError(f'{key}: official clinical + participant CDE counts do not equal total')
        forms = []
        sections = {}
        for field in fields:
            if field['form'] not in forms:
                forms.append(field['form'])
            sections.setdefault(field['form'], [])
            if field['section'] not in sections[field['form']]:
                sections[field['form']].append(field['section'])
        response_rows = sum(str(field['type']).lower() in INPUT_TYPES for field in fields)
        required_response_rows = sum(field['required'] and str(field['type']).lower() in INPUT_TYPES for field in fields)
        classes[key] = {
            **count_meta,
            'filename': source['filename'],
            'download_url': source['url'],
            'dictionaryRows': len(fields),
            'responseRows': response_rows,
            'requiredResponseRows': required_response_rows,
            'requiredDictionaryRows': sum(field['required'] for field in fields),
            'identifierRows': sum(field['identifier'] for field in fields),
            'branchingRows': sum(bool(field['branching']) for field in fields),
            'calculatedRows': sum(field['type'] == 'calc' for field in fields),
            'calculationExpressionRows': sum(field['type'] == 'calc' and bool(field.get('choices')) for field in fields),
            'descriptiveRows': sum(field['type'] == 'descriptive' for field in fields),
            'forms': forms,
            'sectionsByForm': sections,
            'fields': fields,
        }
    files = {
        key: {
            'filename': value['filename'], 'download_url': value['url'], 'sha256': value['sha256'],
            'bytes': value['bytes'], 'dictionaryRows': value['expected_rows']
        } for key, value in cfg['sources'].items()
    }
    for key, value in cfg.get('crosscheck_sources', {}).items():
        files[key] = {'filename': value['filename'], 'download_url': value['url'], 'sha256': value['sha256'], 'bytes': value['bytes']}
    return {
        'schemaVersion': cfg['schema_version'],
        'generatedAt': cfg['schema_generated_at'],
        'source': {
            'title': official['source_title'],
            'ccrfVersion': cfg['ccrf_version'],
            'officialPage': cfg['official_page'],
            'officialCountSource': official['source_url'],
            'officialCountAccessedDate': official['accessed_date'],
            'countDefinition': official['definition'],
            'ccrfPaper': 'https://doi.org/10.3389/fgwh.2026.1833913',
            'files': files,
            'notice': 'Official CDE totals are taken from the Lipedema Foundation CCRF webpage. REDCap dictionary-row counts are reported separately because the workbooks also contain technical, administrative, instructional and calculated rows. Calculated fields are preserved as supplied; no calculation is inferred when the official expression is blank.',
            'calculationPolicy': 'Execute only an exact official CCRF calculation expression after its REDCap behaviour is validated. The current official dictionaries mark calculated fields but provide blank calculation expressions, so LipoScope does not calculate those values.',
        },
        'classifications': classes,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument('--input-dir')
    parser.add_argument('--snapshot', default=str(DATA / 'ccrf_source_snapshot.tsv'))
    parser.add_argument('--output', default=str(BUILD / 'ccrf_schema.rebuilt.json'))
    args = parser.parse_args()
    cfg = load_sources()
    official = load_official_counts()
    if args.input_dir:
        source_dir = Path(args.input_dir)
        groups = {key: parse_xlsx(source_dir / source['filename']) for key, source in cfg['sources'].items()}
    else:
        groups = parse_snapshot(Path(args.snapshot))
    schema = build_schema(groups, cfg, official)
    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(schema, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f"Wrote {out} with " + ', '.join(f'{key}={value["official_cdes"]} official CDEs / {len(value["fields"])} dictionary rows' for key, value in schema['classifications'].items()))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
