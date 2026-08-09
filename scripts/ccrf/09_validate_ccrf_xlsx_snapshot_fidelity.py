#!/usr/bin/env python3
"""
LipoScope v1.0.0 — official CCRF XLSX ↔ frozen TSV fidelity test.

Question:
Can data/ccrf_source_snapshot.tsv be independently and deterministically
regenerated from, and shown to be content-equivalent to, the pinned official
Lipedema Foundation CCRF Version 1 REDCap XLSX dictionaries?


PASS requires:
1. All four XLSX files match the release-pinned byte sizes and SHA-256 hashes.
2. Populated REDCap rows ("Variable / Field Name" present) match the expected
   counts and source-row order.
3. A TSV regenerated directly from those XLSX files is byte-for-byte identical
   to the committed frozen TSV.
4. Every TSV column matches row-by-row, including questions, labels, coded
   choices/calculation cells, notes, validation, bounds, identifiers, branching,
   required status, alignment, survey/matrix metadata, annotations, source-row
   positions, classification and source provenance.
5. The frozen TSV matches the locked release SHA-256 and Git blob SHA-1.


A PASS demonstrates deterministic reproducibility and content-equivalence for
the locked release.

Dependency: openpyxl
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


RELEASE_COMMIT = "42e2077f9c437434ee16749a98316c67a7cbe113"

EXPECTED_SNAPSHOT_SHA256 = (
    "74934b25105072e33e949e2f145eee2e09435028b5cdd528424a55c5de79046b"
)
EXPECTED_SNAPSHOT_GIT_BLOB_SHA1 = (
    "333218105eb82ca676c4fdda540c53d6a61a54c1"
)

SOURCES = {
    "exploratory": {
        "filename": "CCRF_Exploratory_DataDictionary.xlsx",
        "sha256": "a9b4817136b70e670ac93f0ff9d10bb36735479f232e382e825db128bc409ac8",
        "bytes": 126959,
        "rows": 739,
    },
    "supplemental": {
        "filename": "CCRF_Supplemental_DataDictionary.xlsx",
        "sha256": "d9b6c8bb4ff07be0e63d4fb82cf91a5d301a80db0a23ce89a7565151f80cea69",
        "bytes": 115829,
        "rows": 650,
    },
    "supplemental_highly_recommended": {
        "filename": "CCRF_Supplemental_Highly_Recommended_DataDictionary.xlsx",
        "sha256": "5bcfd2130ffd6e27fe92e5dd561b440c710011cdcc37fa65d598b56c71037255",
        "bytes": 105521,
        "rows": 445,
    },
    "core": {
        "filename": "CCRF_Core_DataDictionary.xlsx",
        "sha256": "d507515eb1bf677f444cad03010ff8dc8dcae82986e5bfc25c9bf94dd0151765",
        "bytes": 94699,
        "rows": 300,
    },
}

CLASSIFICATION_ORDER = [
    "exploratory",
    "supplemental",
    "supplemental_highly_recommended",
    "core",
]

# Exact REDCap column names represented in the frozen TSV.
SOURCE_COLUMNS = [
    "Variable / Field Name",
    "Form Name",
    "Section Header",
    "Field Type",
    "Field Label",
    "Choices, Calculations, OR Slider Labels",
    "Field Note",
    "Text Validation Type OR Show Slider Number",
    "Text Validation Min",
    "Text Validation Max",
    "Identifier?",
    "Branching Logic (Show field only if...)",
    "Required Field?",
    "Custom Alignment",
    "Question Number (surveys only)",
    "Matrix Group Name",
    "Matrix Ranking?",
    "Field Annotation",
]

SNAPSHOT_COLUMNS = [
    "classification",
    "source_filename",
    "source_sha256",
    "source_row",
    *SOURCE_COLUMNS,
]


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def git_blob_sha1(data: bytes) -> str:
    header = f"blob {len(data)}\0".encode("utf-8")
    return hashlib.sha1(header + data).hexdigest()


def text(value: Any) -> str:
    """Match the release snapshot's source-cell string normalisation."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value).strip()


def yes_flag(value: Any) -> str:
    return "y" if str(value or "").strip().casefold() in {
        "y", "yes", "true", "1"
    } else ""


def header_index(headers: list[Any]) -> dict[str, int]:
    observed = {str(v or "").strip(): i for i, v in enumerate(headers)}
    missing = [name for name in SOURCE_COLUMNS if name not in observed]
    if missing:
        raise ValueError(
            f"Missing expected REDCap columns: {missing}; "
            f"observed={list(observed)}"
        )
    return {name: observed[name] for name in SOURCE_COLUMNS}


def regenerate_classification(
    classification: str,
    xlsx_path: Path,
    source_lock: dict[str, Any],
) -> tuple[list[dict[str, str]], dict[str, Any]]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=False)
    ws = wb[wb.sheetnames[0]]
    iterator = ws.iter_rows(values_only=True)

    headers = list(next(iterator))
    idx = header_index(headers)

    out: list[dict[str, str]] = []
    ignored_blank_rows: list[int] = []

    for source_row, values in enumerate(iterator, start=2):
        def get(column: str) -> Any:
            i = idx[column]
            return values[i] if i < len(values) else None

        # This is the same operative row rule used by 03_build_ccrf_schema.py.
        if not text(get("Variable / Field Name")):
            ignored_blank_rows.append(source_row)
            continue

        row = {
            "classification": classification,
            "source_filename": source_lock["filename"],
            "source_sha256": source_lock["sha256"],
            "source_row": str(source_row),
        }
        for column in SOURCE_COLUMNS:
            value = get(column)
            if column in {"Identifier?", "Required Field?"}:
                row[column] = yes_flag(value)
            else:
                row[column] = text(value)
        out.append(row)

    populated_source_rows = [int(row["source_row"]) for row in out]
    expected_source_rows = list(range(2, source_lock["rows"] + 2))

    meta = {
        "sheet_name": ws.title,
        "worksheet_max_row": ws.max_row,
        "populated_redcap_rows": len(out),
        "expected_populated_redcap_rows": source_lock["rows"],
        "populated_source_rows_contiguous": (
            populated_source_rows == expected_source_rows
        ),
        "first_populated_source_row": (
            populated_source_rows[0] if populated_source_rows else None
        ),
        "last_populated_source_row": (
            populated_source_rows[-1] if populated_source_rows else None
        ),
        "ignored_blank_or_formatted_rows": len(ignored_blank_rows),
        "first_ignored_blank_row": (
            ignored_blank_rows[0] if ignored_blank_rows else None
        ),
        "last_ignored_blank_row": (
            ignored_blank_rows[-1] if ignored_blank_rows else None
        ),
    }
    return out, meta


def serialize_tsv(rows: list[dict[str, str]]) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(
        buffer,
        fieldnames=SNAPSHOT_COLUMNS,
        delimiter="\t",
        lineterminator="\n",
        extrasaction="raise",
    )
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        if reader.fieldnames != SNAPSHOT_COLUMNS:
            raise ValueError(
                "Frozen TSV columns differ from the expected release columns."
            )
        return list(reader)


def validate(input_dir: Path, snapshot: Path) -> tuple[dict[str, Any], bytes]:
    source_report: dict[str, Any] = {}
    regenerated_rows: list[dict[str, str]] = []

    # 1. Lock each input to the exact pinned official XLSX.
    for key in CLASSIFICATION_ORDER:
        lock = SOURCES[key]
        path = input_dir / lock["filename"]
        if not path.exists():
            source_report[key] = {
                "exists": False,
                "path": str(path),
                "passed": False,
            }
            continue

        raw = path.read_bytes()
        rows, meta = regenerate_classification(key, path, lock)
        regenerated_rows.extend(rows)

        item = {
            "exists": True,
            "path": str(path),
            "filename": lock["filename"],
            "expected_sha256": lock["sha256"],
            "observed_sha256": sha256(raw),
            "sha256_matches_pinned_source": sha256(raw) == lock["sha256"],
            "expected_bytes": lock["bytes"],
            "observed_bytes": len(raw),
            "byte_size_matches_pinned_source": len(raw) == lock["bytes"],
            **meta,
        }
        item["row_count_matches"] = (
            item["populated_redcap_rows"]
            == item["expected_populated_redcap_rows"]
        )
        item["passed"] = all([
            item["sha256_matches_pinned_source"],
            item["byte_size_matches_pinned_source"],
            item["row_count_matches"],
            item["populated_source_rows_contiguous"],
        ])
        source_report[key] = item

    # 2. Lock the frozen TSV to the reviewed release.
    snapshot_bytes = snapshot.read_bytes()
    snapshot_lock = {
        "path": str(snapshot),
        "bytes": len(snapshot_bytes),
        "sha256": sha256(snapshot_bytes),
        "expected_sha256": EXPECTED_SNAPSHOT_SHA256,
        "sha256_matches_release_lock": (
            sha256(snapshot_bytes) == EXPECTED_SNAPSHOT_SHA256
        ),
        "git_blob_sha1": git_blob_sha1(snapshot_bytes),
        "expected_git_blob_sha1": EXPECTED_SNAPSHOT_GIT_BLOB_SHA1,
        "git_blob_matches_release_commit": (
            git_blob_sha1(snapshot_bytes) == EXPECTED_SNAPSHOT_GIT_BLOB_SHA1
        ),
    }

    # 3. Deterministically regenerate the full frozen TSV.
    regenerated_bytes = serialize_tsv(regenerated_rows)
    regeneration = {
        "rows": len(regenerated_rows),
        "bytes": len(regenerated_bytes),
        "sha256": sha256(regenerated_bytes),
        "byte_identical_to_committed_snapshot": (
            regenerated_bytes == snapshot_bytes
        ),
        "same_sha256": sha256(regenerated_bytes) == sha256(snapshot_bytes),
    }

    # 4. Explicit row × column comparison for a reviewer-readable audit.
    frozen_rows = read_tsv(snapshot)
    mismatches: list[dict[str, Any]] = []
    field_differences = 0

    if len(regenerated_rows) != len(frozen_rows):
        mismatches.append({
            "kind": "row_count",
            "regenerated_rows": len(regenerated_rows),
            "frozen_rows": len(frozen_rows),
        })

    for i, (xlsx_row, tsv_row) in enumerate(
        zip(regenerated_rows, frozen_rows), start=1
    ):
        diffs = {}
        for column in SNAPSHOT_COLUMNS:
            if xlsx_row[column] != tsv_row[column]:
                diffs[column] = {
                    "xlsx": xlsx_row[column],
                    "frozen_tsv": tsv_row[column],
                }
                field_differences += 1
        if diffs:
            mismatches.append({
                "snapshot_data_row": i,
                "classification": xlsx_row["classification"],
                "source_row": xlsx_row["source_row"],
                "variable": xlsx_row["Variable / Field Name"],
                "differences": diffs,
            })

    comparison = {
        "columns_compared": SNAPSHOT_COLUMNS,
        "source_metadata_columns_compared": SOURCE_COLUMNS,
        "expected_cumulative_technical_rows": sum(
            SOURCES[k]["rows"] for k in CLASSIFICATION_ORDER
        ),
        "rows_compared": min(len(regenerated_rows), len(frozen_rows)),
        "rows_with_differences": len(mismatches),
        "field_differences": field_differences,
        "mismatch_details": mismatches,
    }

    sources_pass = all(
        source_report.get(k, {}).get("passed", False)
        for k in CLASSIFICATION_ORDER
    )
    expected_total = comparison["expected_cumulative_technical_rows"]

    passed = all([
        sources_pass,
        snapshot_lock["sha256_matches_release_lock"],
        snapshot_lock["git_blob_matches_release_commit"],
        regeneration["rows"] == expected_total,
        regeneration["byte_identical_to_committed_snapshot"],
        regeneration["same_sha256"],
        len(frozen_rows) == expected_total,
        comparison["rows_compared"] == expected_total,
        comparison["rows_with_differences"] == 0,
        comparison["field_differences"] == 0,
    ])

    report = {
        "test": "LipoScope CCRF official XLSX ↔ frozen TSV fidelity",
        "release_commit": RELEASE_COMMIT,
        "question": (
            "Can the frozen TSV be independently and deterministically "
            "regenerated from, and shown to be content-equivalent to, "
            "the pinned official XLSX sources?"
        ),
        "passed": passed,
        "source_files": source_report,
        "committed_snapshot": snapshot_lock,
        "deterministic_regeneration": regeneration,
        "row_and_field_comparison": comparison,
        "interpretation": (
            "PASS: the four exact pinned official XLSX files independently "
            "regenerated the committed frozen TSV byte-for-byte. All 2,134 "
            "cumulative technical rows and every retained source/provenance "
            "column matched with zero differences."
            if passed else
            "FAIL: at least one source-lock, byte-level, row-level, ordering, "
            "provenance or retained-metadata check differed."
        ),
        "worksheet_dimension_note": (
            "Trailing blank/formatted XLSX rows are not CCRF dictionary "
            "records. The validator follows the same operative rule as "
            "03_build_ccrf_schema.py and processes only rows with a populated "
            "'Variable / Field Name'."
        ),
        "provenance_boundary": (
            "This demonstrates deterministic derivability and content-"
            "equivalence for the locked release."
        ),
    }
    return report, regenerated_bytes


def print_summary(report: dict[str, Any]) -> None:
    print("=" * 78)
    print("LipoScope CCRF official XLSX ↔ frozen TSV fidelity test")
    print("=" * 78)

    for key in CLASSIFICATION_ORDER:
        item = report["source_files"][key]
        print(
            f"{'PASS' if item.get('passed') else 'FAIL'}  {key}: "
            f"{item.get('populated_redcap_rows', 'NA')}/"
            f"{item.get('expected_populated_redcap_rows', 'NA')} rows; "
            f"SHA-256 match={item.get('sha256_matches_pinned_source', False)}"
        )

    regen = report["deterministic_regeneration"]
    comp = report["row_and_field_comparison"]
    snap = report["committed_snapshot"]

    print()
    print(f"Snapshot SHA-256 release lock: {snap['sha256_matches_release_lock']}")
    print(f"Snapshot Git blob release lock: {snap['git_blob_matches_release_commit']}")
    print(f"Regenerated rows: {regen['rows']}")
    print(f"Byte-for-byte identical TSV: {regen['byte_identical_to_committed_snapshot']}")
    print(f"Identical TSV SHA-256: {regen['same_sha256']}")
    print(f"Rows compared: {comp['rows_compared']}")
    print(f"Rows with differences: {comp['rows_with_differences']}")
    print(f"Field differences: {comp['field_differences']}")
    print()
    print("OVERALL:", "PASS" if report["passed"] else "FAIL")
    print(report["interpretation"])


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--snapshot", required=True)
    parser.add_argument("--report", default="")
    parser.add_argument("--write-regenerated", default="")
    args = parser.parse_args()

    report, regenerated = validate(
        Path(args.input_dir),
        Path(args.snapshot),
    )
    print_summary(report)

    if args.report:
        Path(args.report).write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
    if args.write_regenerated:
        Path(args.write_regenerated).write_bytes(regenerated)

    return 0 if report["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
